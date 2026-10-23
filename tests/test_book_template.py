"""Word template for two-line books: A4 portrait cut into two A5 landscape pages."""
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.shared import Mm

from page_calibration import DEFAULT_FORMAT as FMT
from tools.make_book_template import (
    COLOR_STYLE, RAISED_STYLE, expanded_pt, line_spacing_pt, make_template, marker_positions_mm,
)

EMU_PER_MM = 36000


def anchors(part):
    """(x_mm, y_mm, size_mm) of every floating picture in a header part."""
    ns = {'wp': 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing'}
    result = []
    for anchor in part._element.findall('.//wp:anchor', ns):
        x = int(anchor.find('wp:positionH/wp:posOffset', ns).text) / EMU_PER_MM
        y = int(anchor.find('wp:positionV/wp:posOffset', ns).text) / EMU_PER_MM
        size = int(anchor.find('wp:extent', ns).get('cx')) / EMU_PER_MM
        result.append((round(x, 2), round(y, 2), round(size, 2)))
    return sorted(result)


class BookTemplateTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        cls.path = make_template(Path(cls.temp.name) / 'book.docx', cell_pitch_mm=FMT.cell_pitch_mm, line_pitch_mm=10)
        cls.doc = Document(str(cls.path))

    @classmethod
    def tearDownClass(cls):
        cls.temp.cleanup()

    def test_expanded_spacing_formula(self):
        self.assertAlmostEqual(expanded_pt(6.21), 0.05, delta=0.01)  # 32.5 pt: the font's cell is nearly the pitch
        self.assertAlmostEqual(expanded_pt(6.2, font_pt=30), 1.37, delta=0.01)

    def test_line_spacing_is_exact_slate_row_pitch(self):
        self.assertAlmostEqual(line_spacing_pt(10), 28.35, delta=0.01)

    def test_marker_positions_mm(self):
        top = {'TL': (10, 10), 'TR': (185, 10), 'BR': (185, 123), 'BL': (10, 123)}
        self.assertEqual(marker_positions_mm(FMT, 0.0), top)
        self.assertEqual(marker_positions_mm(FMT, 148.5)['BR'], (185, 271.5))

    def test_front_header_floats_eight_markers_at_their_positions(self):
        expected = sorted((float(x), float(y), 15.0)
                          for half in (0.0, 148.5) for x, y in marker_positions_mm(FMT, half).values())
        self.assertEqual(anchors(self.doc.sections[0].header), expected)

    def test_back_page_has_no_markers(self):
        back = self.doc.sections[1]
        self.assertFalse(back.header.is_linked_to_previous)
        self.assertEqual(anchors(back.header), [])

    def test_page_is_a4_portrait(self):
        for section in self.doc.sections:
            self.assertEqual((round(section.page_width.mm), round(section.page_height.mm)), (210, 297))

    def test_front_alternates_colored_and_raised_lines_on_both_halves(self):
        styles = [p.style.name for p in self.doc.sections[0].iter_inner_content() if hasattr(p, 'style')]
        self.assertEqual(styles, [COLOR_STYLE, RAISED_STYLE] * (2 * FMT.line_pairs))

    def test_colored_style_uses_braille_font_and_slate_spacing(self):
        style = self.doc.styles[COLOR_STYLE]
        self.assertEqual(style.font.name, 'SOV_braille_only')
        self.assertEqual(style.font.size.pt, 32.5)
        spacing = style.element.rPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}spacing')
        self.assertEqual(int(spacing.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val')), 1)  # 0.05 pt
        self.assertAlmostEqual(style.paragraph_format.line_spacing.pt, 28.35, delta=0.05)

    def test_colored_style_is_lowered_into_the_exact_line(self):
        # 32.5 pt dots are taller than a 10 mm exact line; Word clips row 1 unless lowered.
        position = self.doc.styles[COLOR_STYLE].element.rPr.find(
            '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}position')
        self.assertEqual(int(position.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val')), -6)  # 3 pt in half-points

    def test_marker_size_is_configurable(self):
        from dataclasses import replace
        fmt = replace(FMT, marker_size_mm=12.0)
        with tempfile.TemporaryDirectory() as temp:
            doc = Document(str(make_template(Path(temp) / 'small.docx', 165 / 27, 10, fmt=fmt)))
        placed = anchors(doc.sections[0].header)
        self.assertEqual({size for _, _, size in placed}, {12.0})
        self.assertIn((188.0, 10.0, 12.0), placed)  # TR: 210 - 10 margin - 12

    def test_back_guides_are_mirrored(self):
        # The slate is written from the back, right to left: the guide sits on the mirrored side.
        guides = [p for p in self.doc.sections[1].iter_inner_content() if p.text.strip()]
        self.assertEqual(len(guides), 2 * FMT.line_pairs)
        front_left = self.doc.sections[0].left_margin
        for guide in guides:
            self.assertEqual(guide.alignment, 2)  # right
            self.assertAlmostEqual(guide.paragraph_format.right_indent.mm + self.doc.sections[1].right_margin.mm,
                                   front_left.mm, delta=0.1)


class BookFromTextTests(unittest.TestCase):
    """Text in, book out: colored lines on the front, every dot to punch on the back."""

    def test_guide_dots_are_mirrored_under_the_colored_line(self):
        from tools.make_book_template import guide_dots_mm
        dots = guide_dots_mm(['ก'], FMT, 10.0)  # ก = dots 1, 2, 4, 5 in the first cell
        self.assertEqual(len(dots), 4)
        sheet, page_x, page_y = zip(*dots)
        self.assertEqual(set(sheet), {0})
        left, right = sorted(set(round(x, 2) for x in page_x))
        self.assertAlmostEqual(right - left, FMT.dot_spacing_mm, delta=0.05)
        self.assertGreater(left, 210 / 2)  # the first cell sits on the right of the back side
        top, middle = sorted(set(round(y, 2) for y in page_y))
        self.assertAlmostEqual(middle - top, FMT.dot_spacing_mm, delta=0.05)
        self.assertGreater(top, 28.0 + 10.0)  # in the raised row, one slate row below the colored line

    def test_two_cell_letter_gets_dots_in_two_cells(self):
        from tools.make_book_template import guide_dots_mm
        one, two = guide_dots_mm(['ก'], FMT, 10.0), guide_dots_mm(['ฃ'], FMT, 10.0)
        width = lambda dots: max(x for _, x, _ in dots) - min(x for _, x, _ in dots)  # noqa: E731
        self.assertAlmostEqual(width(one), FMT.dot_spacing_mm, delta=0.05)  # one cell: two columns
        self.assertAlmostEqual(width(two), FMT.cell_pitch_mm, delta=0.05)  # ฃ = [3,5,6] + [1,3]: next cell's left column
        self.assertEqual(len(two), 5)

    def test_cells_follow_the_slate_pitch(self):
        from tools.make_book_template import guide_dots_mm
        dots = guide_dots_mm(['ฮ' * 24], FMT, 10.0)
        xs = sorted({round(x, 3) for _, x, _ in dots}, reverse=True)
        self.assertAlmostEqual((xs[0] - xs[-2]) / 23, FMT.cell_pitch_mm, delta=0.02)

    def test_line_longer_than_the_slate_row_is_rejected(self):
        with tempfile.TemporaryDirectory() as temp, self.assertRaises(ValueError) as caught:
            make_template(Path(temp) / 'b.docx', FMT.cell_pitch_mm, 10, lines=['ก', 'ฮ' * 25, 'ไก่', 'ฮ' * 26])
        message = str(caught.exception)
        self.assertIn('25 cells: ' + 'ฮ' * 25, message)  # every long line, with its text
        self.assertIn('26 cells: ' + 'ฮ' * 26, message)
        self.assertNotIn('ไก่', message)

    def test_nine_lines_make_two_sheets_with_markers_on_each_front(self):
        lines = ['ไก่'] * 9
        with tempfile.TemporaryDirectory() as temp:
            path = make_template(Path(temp) / 'b.docx', FMT.cell_pitch_mm, 10, lines=lines)
            doc = Document(str(path))
            answers = sorted(p.name for p in Path(temp).glob('b_page*.txt'))
            page3 = (Path(temp) / 'b_page3.txt').read_text(encoding='utf-8')
        self.assertEqual(len(doc.sections), 4)  # front, back, front, back
        self.assertEqual(len(anchors(doc.sections[0].header)), 8)
        self.assertEqual(len(anchors(doc.sections[2].header)), 8)
        self.assertEqual(answers, ['b_page1.txt', 'b_page2.txt', 'b_page3.txt'])
        self.assertEqual(page3.split(), ['ไก่'])

    def test_front_lines_hold_the_text(self):
        with tempfile.TemporaryDirectory() as temp:
            doc = Document(str(make_template(Path(temp) / 'b.docx', FMT.cell_pitch_mm, 10, lines=['ไก่ มา', 'ดี'])))
        colored = [p.text for p in doc.sections[0].iter_inner_content() if p.style.name == COLOR_STYLE]
        self.assertEqual(colored[:2], ['ไก่ มา', 'ดี'])
        self.assertEqual(len(colored), 2 * FMT.line_pairs)

    def test_wide_letters_need_no_extra_spacing(self):
        # The font's cell is the slate pitch: a 2-cell letter is already two cells wide.
        w = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
        with tempfile.TemporaryDirectory() as temp:
            doc = Document(str(make_template(Path(temp) / 'b.docx', FMT.cell_pitch_mm, 10, lines=['ก ฃ ค'])))
        paragraph = next(p for p in doc.paragraphs if p.text == 'ก ฃ ค')
        spacing = {}
        for run in paragraph.runs:
            node = run._r.rPr.find(w + 'spacing') if run._r.rPr is not None else None
            spacing[run.text] = None if node is None else int(node.get(w + 'val'))
        self.assertEqual(spacing, {'ก ฃ ค': None})  # one run in the style's spacing

    def test_back_side_carries_the_guide_picture(self):
        with tempfile.TemporaryDirectory() as temp:
            doc = Document(str(make_template(Path(temp) / 'b.docx', FMT.cell_pitch_mm, 10, lines=['ไก่'])))
        self.assertEqual(anchors(doc.sections[1].header), [(0.0, 0.0, 210.0)])


class PitchTestPageTests(unittest.TestCase):
    """Lay the slate over the page: the line whose 24 full cells sit in the windows wins."""

    def test_one_full_cell_line_per_expanded_value(self):
        from tools.make_book_template import FULL_CELL, make_pitch_test
        values = (0.6, 0.8, 1.0, 1.1, 1.2, 1.4)
        with tempfile.TemporaryDirectory() as temp:
            doc = Document(str(make_pitch_test(Path(temp) / 'pitch.docx', values)))
        w = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
        lines = [p for p in doc.paragraphs if FULL_CELL in p.text]
        self.assertEqual(len(lines), len(values))
        for paragraph, value in zip(lines, values):
            cells = paragraph.runs[0]
            self.assertEqual(cells.text, FULL_CELL * 24)
            self.assertEqual(paragraph.style.name, COLOR_STYLE)
            self.assertEqual(int(cells._r.rPr.find(w + 'spacing').get(w + 'val')), round(value * 20))
            self.assertIn(f'{value:.1f} pt', paragraph.text)  # label after the cells

    def test_full_cell_is_the_six_dot_glyph(self):
        from tools.make_book_template import FULL_CELL
        self.assertEqual(FULL_CELL, 'ฮ')  # SOV_braille_only draws ฮ as all six dots, one cell wide


if __name__ == '__main__':
    unittest.main()
