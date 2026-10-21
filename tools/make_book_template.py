"""Write the Word file for two-line books (A4 portrait, cut into two A5 pages).

Front: 4 ArUco markers per A5 half, and alternating lines styled for the colored
Braille line and the raised line (left empty for the slate).
Back (print double-sided, flip on the long edge), mirrored because the slate is
written from the back: with --text, a grey dot wherever the slate must punch;
without it, a grey mark where each raised line starts.

Usage: make_book_template.py --text book.txt --out output/book.docx
       make_book_template.py --out output/book_template.docx   (blank template)
"""

from __future__ import annotations

import argparse
import copy
import sys
import tempfile
from pathlib import Path

import cv2
import numpy as np
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Mm, Pt, RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from generate_test import text_to_braille_cells  # noqa: E402
from page_calibration import CORNERS, DEFAULT_FORMAT, PageFormat, page_format  # noqa: E402
from tools.export_markers import export_markers  # noqa: E402

FONT = 'SOV_braille_only'
FONT_PT = 30
FULL_CELL = 'ฮ'  # SOV_braille_only draws it as all six dots, one cell wide
PITCH_TEST_VALUES = (0.6, 0.8, 1.0, 1.1, 1.2, 1.4, 1.6)
LOWERED_PT = 2  # the exact 10 mm line is shorter than 30 pt: lower the dots so row 1 is not clipped
PT_MM = 25.4 / 72
COLOR_STYLE = 'Braille สี'
RAISED_STYLE = 'บรรทัดนูน'
HALF_MM = 148.5  # half of a portrait A4: where the sheet is cut
TOP_MM = 28.0  # first colored line of each half
LEFT_MM, RIGHT_MM = 30.0, 20.0
DOT_BLUE = RGBColor(0x00, 0x50, 0xDC)  # hue inside the reader's blue range
GUIDE = '◄'


def expanded_pt(cell_pitch_mm: float, font_pt: float = FONT_PT) -> float:
    """Word 'Expanded by' spacing that makes the font's cells match the slate pitch."""
    return cell_pitch_mm / PT_MM - 0.54 * font_pt  # 0.54 em: SOV_braille_only advance width


def line_spacing_pt(line_pitch_mm: float) -> float:
    return line_pitch_mm / PT_MM


def marker_positions_mm(fmt: PageFormat, top_mm: float) -> dict[str, tuple[float, float]]:
    """Top-left corner of each marker on the A4 sheet, for the A5 half starting at top_mm."""
    corners = fmt.marker_corners_mm()
    return {name: (float(corners[mid][0][0]), float(corners[mid][0][1]) + top_mm)
            for name, mid in zip(CORNERS, fmt.marker_ids)}


def guide_dots_mm(lines: list[str], fmt: PageFormat, line_pitch_mm: float,
                  cell_pitch_mm: float | None = None) -> list[tuple[int, float, float]]:
    """(sheet, x, y) on the BACK side of every dot the slate must punch, in A4 page mm.

    Cells come from the project's Thai encoder, which matches SOV_braille_only
    cell for cell (2-cell letters such as ฃ included). The back is the front
    mirrored left to right, and the raised row is one slate row below the colored line.
    """
    pitch = _printed_pitch_mm(cell_pitch_mm or fmt.cell_pitch_mm)
    rows_mm = _dot_rows_in_line_mm(line_pitch_mm)
    per_half, dots = fmt.line_pairs, []
    for index, text in enumerate(lines):
        sheet, slot = divmod(index, 2 * per_half)
        half, pair = divmod(slot, per_half)
        raised_top = half * HALF_MM + TOP_MM + (2 * pair + 1) * line_pitch_mm
        for k, cell in enumerate(text_to_braille_cells(text, 'thai')):
            for dot in cell or ():
                x_front = LEFT_MM + k * pitch + _DOT_COLUMNS_MM[(dot - 1) // 3]
                dots.append((sheet, 210.0 - x_front, raised_top + rows_mm[(dot - 1) % 3]))
    return dots


def make_template(out: Path, cell_pitch_mm: float, line_pitch_mm: float,
                  fmt: PageFormat = DEFAULT_FORMAT, lines: list[str] | None = None) -> Path:
    """Blank template (``lines`` None), or a book: text on the front, punch guide on the back.

    A book also writes one answer file per A5 page (``<out>_pageN.txt``) for
    tools/diagnostics/page_map_accuracy.py.
    """
    per_sheet = 2 * fmt.line_pairs  # colored lines per A4 sheet (two A5 halves)
    if lines is not None:
        too_long = [f'  {cells} cells: {text}' for text in lines
                    if (cells := len(text_to_braille_cells(text, 'thai'))) > fmt.cells_per_line]
        if too_long:
            raise ValueError(f'a slate row holds {fmt.cells_per_line} cells; split these lines:\n' + '\n'.join(too_long))
    doc = Document()
    _add_styles(doc, cell_pitch_mm, line_pitch_mm)
    sheets = [lines[i:i + per_sheet] for i in range(0, len(lines), per_sheet)] if lines else [None]
    rows = 2 * fmt.line_pairs  # per A5 half: colored, raised, colored, raised, ...
    gap_mm = HALF_MM - rows * line_pitch_mm  # the second half starts one A5 page lower
    for number, sheet in enumerate(sheets):
        front = _side(doc, 2 * rows, gap_mm, rows, last=False)
        for index, paragraph in enumerate(front):
            paragraph.style = doc.styles[COLOR_STYLE if index % 2 == 0 else RAISED_STYLE]
            if sheet is not None and index % 2 == 0 and index // 2 < len(sheet):
                _add_braille_text(paragraph, sheet[index // 2], cell_pitch_mm)
        back = _side(doc, 2 * rows, gap_mm, rows, last=number == len(sheets) - 1)
        for index, paragraph in enumerate(back):
            paragraph.style = doc.styles[RAISED_STYLE]
            if sheet is None and index % 2 == 1:  # blank template: mark where the slate's first cell goes
                paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                paragraph.paragraph_format.right_indent = Mm(LEFT_MM - RIGHT_MM)
                run = paragraph.add_run(GUIDE)
                run.font.size, run.font.color.rgb = Pt(10), RGBColor(0xB0, 0xB0, 0xB0)
    for section in doc.sections:
        _set_page(section, top_mm=TOP_MM)
        section.header.is_linked_to_previous = False
    with tempfile.TemporaryDirectory() as temp:
        paths = dict(zip(CORNERS, export_markers(Path(temp), px=600, fmt=fmt)))
        guides = _guide_pictures(Path(temp), lines, fmt, line_pitch_mm, cell_pitch_mm) if lines else {}
        for number, section in enumerate(doc.sections):
            header = section.header.paragraphs[0]
            if number % 2 == 0:
                for top in (0.0, HALF_MM):
                    for name, (x, y) in marker_positions_mm(fmt, top).items():
                        _float_picture(header.add_run(), paths[name], x, y, fmt.marker_size_mm)
            elif number // 2 in guides:
                _float_picture(header.add_run(), guides[number // 2], 0.0, 0.0, 210.0, 297.0)
        out.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(out))
    if lines:
        half = fmt.line_pairs
        for page, start in enumerate(range(0, len(lines), half), 1):
            answer = out.with_name(f'{out.stem}_page{page}.txt')
            answer.write_text('\n'.join(lines[start:start + half]) + '\n', encoding='utf-8')
    return out


def _add_braille_text(paragraph, text: str, cell_pitch_mm: float) -> None:
    """Add the text so every Braille cell lands on the slate pitch.

    Word adds 'Expanded' once per character, but a wide letter (ฃ = 2 cells)
    needs it once per cell. Runs of one-cell characters keep the style's spacing.
    """
    expanded = expanded_pt(cell_pitch_mm)
    plain = ''
    for char in text:
        cells = len(text_to_braille_cells(char, 'thai')) if char.strip() else 1
        if cells == 1:
            plain += char
            continue
        if plain:
            paragraph.add_run(plain)
            plain = ''
        run = paragraph.add_run(char)
        run._r.get_or_add_rPr().append(parse_xml(f'<w:spacing {nsdecls("w")} w:val="{round(cells * expanded * 20)}"/>'))
    if plain:
        paragraph.add_run(plain)


def _side(doc, count: int, gap_mm: float, rows: int, last: bool) -> list:
    """``count`` paragraphs for one printed side; the second half starts lower.

    Unless ``last``, the side ends with a section break, which python-docx puts
    in a new paragraph: that paragraph is the side's final line.
    """
    paragraphs = [doc.add_paragraph() for _ in range(count - (0 if last else 1))]
    if not last:
        doc.add_section()
        paragraphs.append(doc.paragraphs[-1])
    paragraphs[rows].paragraph_format.space_before = Mm(gap_mm)
    return paragraphs


def _guide_pictures(folder: Path, lines, fmt, line_pitch_mm, cell_pitch_mm, px_per_mm: int = 12) -> dict[int, Path]:
    """One transparent PNG per back side: grey dots where the slate must punch."""
    radius = round(fmt.dot_diameter_mm / 2 * px_per_mm)
    pictures: dict[int, np.ndarray] = {}
    for sheet, x, y in guide_dots_mm(lines, fmt, line_pitch_mm, cell_pitch_mm):
        canvas = pictures.setdefault(sheet, np.zeros((297 * px_per_mm, 210 * px_per_mm, 4), np.uint8))
        cv2.circle(canvas, (round(x * px_per_mm), round(y * px_per_mm)), radius, (150, 150, 150, 255), -1, cv2.LINE_AA)
    paths = {}
    for sheet, canvas in pictures.items():
        paths[sheet] = folder / f'guide_{sheet}.png'
        if not cv2.imwrite(str(paths[sheet]), canvas):
            raise OSError(f'cannot write {paths[sheet]}')
    return paths


def _printed_pitch_mm(cell_pitch_mm: float) -> float:
    """The pitch Word lays out: 'Expanded' is stored in twentieths of a point."""
    return (0.54 * FONT_PT + round(expanded_pt(cell_pitch_mm) * 20) / 20) * PT_MM


def _dot_rows_in_line_mm(line_pitch_mm: float) -> tuple[float, float, float]:
    """Dot rows 1-3 below the top of an exact-height line.

    shortcut: Word puts the baseline at ~80% of an exact line (ascent 800/1000);
    rows sit 720/500/280 font units above it, lowered by LOWERED_PT. Re-measure
    on a printed page if the slate rows do not cover the grey dots.
    """
    baseline = 0.8 * line_pitch_mm + LOWERED_PT * PT_MM
    em_mm = FONT_PT * PT_MM
    return tuple(baseline - units / 1000 * em_mm for units in (720, 500, 280))


_DOT_COLUMNS_MM = (0.14 * FONT_PT * PT_MM, 0.36 * FONT_PT * PT_MM)  # glyph columns at 140/360 font units


def make_pitch_test(out: Path, values=PITCH_TEST_VALUES, line_pitch_mm: float = 10.0) -> Path:
    """One line of 24 full cells per 'Expanded' value, a raised-line gap between them.

    Print at 100%, lay the slate on each line: pick the value whose 24 cells all
    sit inside the slate windows, then build the template with that pitch.
    """
    doc = Document()
    _set_page(doc.sections[0], top_mm=TOP_MM)
    _add_styles(doc, DEFAULT_FORMAT.cell_pitch_mm, line_pitch_mm)
    for value in values:
        line = doc.add_paragraph(style=doc.styles[COLOR_STYLE])
        cells = line.add_run(FULL_CELL * 24)
        cells._r.get_or_add_rPr().append(parse_xml(f'<w:spacing {nsdecls("w")} w:val="{round(value * 20)}"/>'))
        label = line.add_run(f'   {value:.1f} pt')
        label.font.name, label.font.size, label.font.color.rgb = 'Tahoma', Pt(9), RGBColor(0x80, 0x80, 0x80)
        label._r.get_or_add_rPr().get_or_add_rFonts().set(qn('w:cs'), 'Tahoma')
        doc.add_paragraph(style=doc.styles[RAISED_STYLE])
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return out


def _set_page(section, top_mm: float) -> None:
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin, section.bottom_margin = Mm(top_mm), Mm(10)
    section.left_margin, section.right_margin = Mm(LEFT_MM), Mm(RIGHT_MM)
    section.header_distance = Mm(5)


def _add_styles(doc, cell_pitch_mm: float, line_pitch_mm: float) -> None:
    for name in (COLOR_STYLE, RAISED_STYLE):
        style = doc.styles.add_style(name, 1)  # paragraph style
        style.base_style = doc.styles['Normal']
        layout = style.paragraph_format
        layout.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        layout.line_spacing = Pt(line_spacing_pt(line_pitch_mm))
        layout.space_before = layout.space_after = Pt(0)
    style = doc.styles[COLOR_STYLE]
    style.font.name, style.font.size, style.font.color.rgb = FONT, Pt(FONT_PT), DOT_BLUE
    rpr = style.element.get_or_add_rPr()
    rpr.rFonts.set(qn('w:cs'), FONT)  # Thai text is complex script in Word
    rpr.append(parse_xml(f'<w:szCs {nsdecls("w")} w:val="{FONT_PT * 2}"/>'))
    rpr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:val="{round(expanded_pt(cell_pitch_mm) * 20)}"/>'))
    rpr.append(parse_xml(f'<w:position {nsdecls("w")} w:val="{-LOWERED_PT * 2}"/>'))  # half-points


def _float_picture(run, path: Path, x_mm: float, y_mm: float, width_mm: float,
                   height_mm: float | None = None) -> None:
    """Insert a picture at a fixed page position (Word: In Front of Text, locked)."""
    run.add_picture(str(path), width=Mm(width_mm), height=Mm(height_mm or width_mm))
    inline = run._r.xpath('.//wp:inline')[-1]
    emu = lambda mm: round(mm * 36000)  # noqa: E731
    anchor = parse_xml(
        f'<wp:anchor {nsdecls("wp", "a", "pic", "r")} distT="0" distB="0" distL="0" distR="0" '
        f'simplePos="0" relativeHeight="251659264" behindDoc="0" locked="1" layoutInCell="1" '
        f'allowOverlap="1"><wp:simplePos x="0" y="0"/>'
        f'<wp:positionH relativeFrom="page"><wp:posOffset>{emu(x_mm)}</wp:posOffset></wp:positionH>'
        f'<wp:positionV relativeFrom="page"><wp:posOffset>{emu(y_mm)}</wp:posOffset></wp:positionV>'
        f'</wp:anchor>')
    anchor.append(copy.deepcopy(inline.find(qn('wp:extent'))))
    anchor.append(parse_xml(f'<wp:effectExtent {nsdecls("wp")} l="0" t="0" r="0" b="0"/>'))
    anchor.append(parse_xml(f'<wp:wrapNone {nsdecls("wp")}/>'))
    for tag in ('wp:docPr', 'wp:cNvGraphicFramePr', 'a:graphic'):
        anchor.append(copy.deepcopy(inline.find(qn(tag))))
    inline.getparent().replace(inline, anchor)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('--cell-pitch-mm', type=float, default=DEFAULT_FORMAT.cell_pitch_mm, help='slate cell pitch')
    parser.add_argument('--line-pitch-mm', type=float, default=10.0, help='slate row pitch')
    parser.add_argument('--marker-size-mm', type=float,
                        help=f'printed marker side (default {DEFAULT_FORMAT.marker_size_mm:g}); '
                             'pass the same value to finger_reader.py')
    parser.add_argument('--out', type=Path, default=Path('output/book_template.docx'))
    parser.add_argument('--pitch-test', action='store_true',
                        help='write a page of full-cell lines at several Expanded values to check against the slate')
    parser.add_argument('--text', type=Path,
                        help='UTF-8 text, one Braille line per text line: builds the book with a punch guide on the back')
    args = parser.parse_args()
    if args.pitch_test:
        print(make_pitch_test(args.out, line_pitch_mm=args.line_pitch_mm))
        return 0
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, 'reconfigure'):
            stream.reconfigure(encoding='utf-8')  # Thai text in messages
    lines = None
    if args.text:
        if not args.text.is_file():
            print(f'ไม่พบไฟล์ข้อความ: {args.text}', file=sys.stderr)
            return 1
        lines = [line.strip() for line in args.text.read_text(encoding='utf-8-sig').splitlines() if line.strip()]
    try:
        path = make_template(args.out, args.cell_pitch_mm, args.line_pitch_mm,
                             fmt=page_format(args.marker_size_mm), lines=lines)
    except (ValueError, PermissionError) as error:  # long lines, or the .docx is open in Word
        print(error, file=sys.stderr)
        return 1
    print(f'{path}  Expanded {expanded_pt(args.cell_pitch_mm):.2f} pt, '
          f'line spacing exactly {line_spacing_pt(args.line_pitch_mm):.2f} pt')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
